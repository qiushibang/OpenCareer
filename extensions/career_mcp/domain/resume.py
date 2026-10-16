"""Local resume extraction and validation primitives."""

from __future__ import annotations

import io
import copy
import re
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Protocol

if TYPE_CHECKING:
    from ..services.resume_ocr import NativePdfPage, ResumeOcrResult

MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_RESUME_CHARS = 200_000
MAX_PDF_PAGES = 20
MIN_RESUME_CHARS = 20
RESUME_SECTIONS = ("basic", "education", "experience", "project", "skill")
_PROPOSAL_KEYS = set(RESUME_SECTIONS) | {"missing_fields", "warnings"}


class ResumeInputError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ExtractedResume:
    filename: str
    source_kind: str
    text: str
    raw_bytes: bytes
    extraction_method: str
    extraction_metadata: dict[str, object]


class ResumeOcrExtractor(Protocol):
    def extract(
        self,
        pages: tuple[NativePdfPage, ...],
        *,
        raw_pdf: bytes,
    ) -> ResumeOcrResult:
        raise NotImplementedError


def normalize_resume_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [
        re.sub(r"[^\S\n]+", " ", line).strip()
        for line in normalized.split("\n")
    ]
    output: list[str] = []
    blank = False
    for line in lines:
        if line:
            output.append(line)
            blank = False
        elif output and not blank:
            output.append("")
            blank = True
    return "\n".join(output).strip()


def extract_resume_text(
    *,
    filename: str,
    mime: str,
    raw_bytes: bytes | None = None,
    text: str | None = None,
    ocr_service: ResumeOcrExtractor | None = None,
) -> ExtractedResume:
    safe_name = Path(filename).name or "resume"
    suffix = Path(safe_name).suffix.lower()
    source_kind = _source_kind(suffix, mime)
    raw = raw_bytes or b""
    if len(raw) > MAX_RESUME_BYTES:
        raise ResumeInputError("FILE_TOO_LARGE", "resume exceeds 10 MB")

    if source_kind == "text":
        extracted = normalize_resume_text(text or "")
        raw = extracted.encode("utf-8")
        extraction_method = "native"
        extraction_metadata = _native_metadata(total_pages=1)
    elif source_kind == "pdf":
        pages = _extract_pdf_pages(raw)
        extracted = normalize_resume_text(
            "\n\n".join(page.text for page in pages)
        )
        extraction_method = "native"
        extraction_metadata = _native_metadata(total_pages=len(pages))
        if ocr_service is not None and any(
            len(page.text) < _min_native_page_characters()
            for page in pages
        ):
            result = _extract_with_ocr(ocr_service, pages, raw)
            extracted = normalize_resume_text(result.text)
            extraction_method = result.extraction_method
            extraction_metadata = copy.deepcopy(result.metadata)
    else:
        extracted = _extract_docx(raw)
        extraction_method = "native"
        extraction_metadata = _native_metadata(total_pages=1)

    if len(extracted) < MIN_RESUME_CHARS:
        raise ResumeInputError(
            "NO_EXTRACTABLE_TEXT" if source_kind != "text" else "INVALID_INPUT",
            "resume contains too little extractable text",
        )
    if len(extracted) > MAX_RESUME_CHARS:
        raise ResumeInputError(
            "INVALID_INPUT",
            "resume text exceeds 200000 characters",
        )
    return ExtractedResume(
        filename=safe_name,
        source_kind=source_kind,
        text=extracted,
        raw_bytes=raw,
        extraction_method=extraction_method,
        extraction_metadata=extraction_metadata,
    )


def validate_resume_proposal(proposal: dict) -> dict:
    if not isinstance(proposal, dict):
        raise ResumeInputError("INVALID_INPUT", "proposal must be an object")
    unknown = set(proposal) - _PROPOSAL_KEYS
    if unknown:
        raise ResumeInputError(
            "INVALID_INPUT",
            f"unknown proposal field: {sorted(unknown)[0]}",
        )
    missing = _PROPOSAL_KEYS - set(proposal)
    if missing:
        raise ResumeInputError(
            "INVALID_INPUT",
            f"missing proposal field: {sorted(missing)[0]}",
        )
    result = copy.deepcopy(proposal)
    if not isinstance(result["basic"], dict):
        raise ResumeInputError("INVALID_INPUT", "basic must be an object")
    for section in RESUME_SECTIONS[1:]:
        if not isinstance(result[section], list):
            raise ResumeInputError(
                "INVALID_INPUT",
                f"{section} must be a list",
            )
    for key in ("missing_fields", "warnings"):
        if not isinstance(result[key], list):
            raise ResumeInputError(
                "INVALID_INPUT",
                f"{key} must be a list",
            )
    _validate_value(result)
    _validate_dates(result)
    return result


def mask_candidate_profile(profile: dict) -> dict:
    masked = copy.deepcopy(profile)
    basic = masked.get("sections", {}).get("basic")
    if not isinstance(basic, dict):
        return masked
    name = basic.get("name")
    if isinstance(name, str) and name:
        parts = name.split()
        basic["name"] = " ".join(
            part[0] + ("*" * (2 if len(part) <= 3 else 3))
            for part in parts
        )
    phone = basic.get("phone")
    if isinstance(phone, str) and phone:
        basic["phone"] = "*" * max(0, len(phone) - 4) + phone[-4:]
    email = basic.get("email")
    if isinstance(email, str) and "@" in email:
        local, domain = email.split("@", 1)
        basic["email"] = (local[:1] or "*") + "***@" + domain
    return masked


def _source_kind(suffix: str, mime: str) -> str:
    if suffix == ".txt" and mime in {"text/plain", ""}:
        return "text"
    if suffix == ".pdf" and mime == "application/pdf":
        return "pdf"
    if suffix == ".docx" and mime == (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ):
        return "docx"
    raise ResumeInputError(
        "UNSUPPORTED_FORMAT",
        "resume must be text, PDF, or DOCX",
    )


def _extract_pdf_pages(raw: bytes) -> tuple[NativePdfPage, ...]:
    try:
        from pypdf import PdfReader
        from ..services.resume_ocr import NativePdfPage

        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted:
            raise ResumeInputError(
                "INVALID_INPUT",
                "password-protected PDFs are not supported",
            )
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ResumeInputError(
                "PDF_PAGE_LIMIT",
                "PDF exceeds 20 pages",
            )
        if not reader.pages:
            raise ResumeInputError("INVALID_INPUT", "PDF contains no pages")
        return tuple(
            NativePdfPage(
                page_number=index,
                text=normalize_resume_text(page.extract_text() or ""),
            )
            for index, page in enumerate(reader.pages, start=1)
        )
    except ResumeInputError:
        raise
    except Exception as exc:
        raise ResumeInputError("INVALID_INPUT", "could not read PDF") from exc


def _extract_with_ocr(
    ocr_service: ResumeOcrExtractor,
    pages: tuple[NativePdfPage, ...],
    raw: bytes,
) -> ResumeOcrResult:
    from ..services.resume_ocr import ResumeOcrError

    try:
        return ocr_service.extract(pages, raw_pdf=raw)
    except ResumeOcrError as exc:
        messages = {
            "OCR_UNAVAILABLE": "local OCR runtime is unavailable",
            "OCR_PAGE_LIMIT": "PDF requires OCR on more than 10 pages",
            "OCR_FAILED": "local OCR could not recover resume text",
            "NO_EXTRACTABLE_TEXT": (
                "resume contains too little extractable text"
            ),
        }
        code = exc.code if exc.code in messages else "OCR_FAILED"
        raise ResumeInputError(code, messages[code]) from exc


def _min_native_page_characters() -> int:
    from ..services.resume_ocr import MIN_NATIVE_PAGE_CHARS

    return MIN_NATIVE_PAGE_CHARS


def _native_metadata(*, total_pages: int) -> dict[str, object]:
    return {
        "version": 1,
        "method": "native",
        "total_pages": total_pages,
        "native_pages": list(range(1, total_pages + 1)),
        "ocr_pages": [],
        "ocr_attempted_pages": [],
        "mean_confidence": None,
        "quality": "native",
        "review_required": False,
        "warning_codes": [],
        "engine": None,
    }


def _extract_docx(raw: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(raw))
        chunks = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        chunks.append(cell.text)
        return normalize_resume_text("\n".join(chunks))
    except Exception as exc:
        raise ResumeInputError("INVALID_INPUT", "could not read DOCX") from exc


def _validate_value(value: object) -> None:
    if isinstance(value, str):
        if len(value) > 10_000:
            raise ResumeInputError("INVALID_INPUT", "proposal string too long")
        return
    if isinstance(value, list):
        if len(value) > 500:
            raise ResumeInputError("INVALID_INPUT", "proposal list too long")
        for item in value:
            _validate_value(item)
        return
    if isinstance(value, dict):
        for key, item in value.items():
            if not isinstance(key, str):
                raise ResumeInputError(
                    "INVALID_INPUT",
                    "proposal keys must be strings",
                )
            _validate_value(item)
        return
    if value is not None and not isinstance(value, (int, float, bool)):
        raise ResumeInputError(
            "INVALID_INPUT",
            "proposal contains unsupported value",
        )


def _validate_dates(proposal: dict) -> None:
    pattern = re.compile(r"^\d{4}(?:-(?:0[1-9]|1[0-2]))?$")
    for section in ("education", "experience", "project"):
        for item in proposal[section]:
            if not isinstance(item, dict):
                raise ResumeInputError(
                    "INVALID_INPUT",
                    f"{section} items must be objects",
                )
            for key in ("start_date", "end_date"):
                value = item.get(key)
                if value in (None, "present"):
                    continue
                if not isinstance(value, str) or not pattern.match(value):
                    raise ResumeInputError(
                        "INVALID_INPUT",
                        f"invalid {key}",
                    )
