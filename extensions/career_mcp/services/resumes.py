"""Application service for private resume staging and confirmed facts."""

from __future__ import annotations

import hashlib
import html
import io
import os
import secrets
from collections.abc import Callable
from datetime import datetime, timedelta, timezone
from pathlib import Path

from ..config import CareerConfig
from ..database.migrations import connect
from ..database.resume_repository import ResumeRepository
from ..domain.resume import (
    RESUME_SECTIONS,
    ResumeInputError,
    extract_resume_text,
    mask_candidate_profile,
    validate_resume_proposal,
)
from .resume_ocr import (
    NativePdfPage,
    PdfiumPageRenderer,
    ResumeOcrResult,
    ResumeOcrService,
)


class ResumeServiceError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class _LazyResumeOcrService:
    def __init__(
        self,
        factory: Callable[[], ResumeOcrService],
    ) -> None:
        self._factory = factory

    def extract(
        self,
        pages: tuple[NativePdfPage, ...],
        *,
        raw_pdf: bytes,
    ) -> ResumeOcrResult:
        return self._factory().extract(pages, raw_pdf=raw_pdf)


class ResumeService:
    def __init__(
        self,
        config: CareerConfig | None = None,
        *,
        private_root: Path | None = None,
        clock: Callable[[], datetime] | None = None,
        ocr_service: ResumeOcrService | None = None,
        ocr_service_factory: Callable[[], ResumeOcrService] | None = None,
    ) -> None:
        if ocr_service is not None and ocr_service_factory is not None:
            raise ValueError(
                "provide either ocr_service or ocr_service_factory"
            )
        self.config = config or CareerConfig.from_env()
        self.private_root = (
            private_root.resolve()
            if private_root is not None
            else self.config.private_resume_dir.parent
        )
        self._clock = clock or (lambda: datetime.now(timezone.utc))
        factory = ocr_service_factory or self._default_ocr_service
        self._ocr_service = (
            ocr_service
            if ocr_service is not None
            else _LazyResumeOcrService(factory)
        )

    def stage_resume(
        self,
        *,
        filename: str,
        mime: str,
        raw_bytes: bytes | None = None,
        text: str | None = None,
    ) -> dict:
        try:
            extracted = extract_resume_text(
                filename=filename,
                mime=mime,
                raw_bytes=raw_bytes,
                text=text,
                ocr_service=self._ocr_service,
            )
        except ResumeInputError as exc:
            raise ResumeServiceError(exc.code, str(exc)) from exc
        content_hash = hashlib.sha256(extracted.raw_bytes).hexdigest()
        suffix = f".{extracted.source_kind}"
        relative_path = f"resumes/{content_hash}{suffix}"
        target = self.private_root / relative_path
        now = self._now()
        expires = now + timedelta(hours=24)
        connection = None
        created_original = False
        try:
            created_original = self._publish_private_original(
                target,
                extracted.raw_bytes,
            )
            connection = connect(self.config.database_path)
            repository = ResumeRepository(connection)
            connection.execute("BEGIN")
            record = repository.create_or_reuse_import(
                content_hash=content_hash,
                upload_token=secrets.token_urlsafe(32),
                source_kind=extracted.source_kind,
                original_filename=extracted.filename,
                private_path=relative_path,
                extracted_text=extracted.text,
                extraction_method=extracted.extraction_method,
                extraction_metadata=extracted.extraction_metadata,
                expires_at=self._iso(expires),
                now=self._iso(now),
            )
            connection.commit()
        except Exception as exc:
            if connection is not None:
                connection.rollback()
            if created_original:
                target.unlink(missing_ok=True)
            raise ResumeServiceError(
                "PERSISTENCE_ERROR",
                "could not store the private resume",
            ) from exc
        finally:
            if connection is not None:
                connection.close()
        return {
            "upload_token": record["upload_token"],
            "source_kind": record["source_kind"],
            "filename": record["original_filename"],
            "sha256": record["content_hash"],
            "characters": len(record["extracted_text"]),
            "expires_at": record["expires_at"],
            **self._public_extraction(record),
        }

    def get_resume_import_text(self, upload_token: str) -> dict:
        record = self._valid_import(upload_token)
        return {
            "upload_token": upload_token,
            "source_kind": record["source_kind"],
            "filename": record["original_filename"],
            "text": record["extracted_text"],
            "characters": len(record["extracted_text"]),
            "extraction_metadata": record["extraction_metadata"],
            **self._public_extraction(record),
        }

    def preview_resume_import(
        self,
        upload_token: str,
        proposal: dict,
    ) -> dict:
        record = self._valid_import(upload_token)
        try:
            normalized = validate_resume_proposal(proposal)
        except ResumeInputError as exc:
            raise ResumeServiceError(exc.code, str(exc)) from exc
        current = self._profile_or_none()
        conflicts = []
        if current is not None:
            for section in RESUME_SECTIONS:
                existing = current["sections"].get(section)
                proposed = normalized[section]
                if existing is not None and existing != proposed:
                    conflicts.append(
                        {
                            "section": section,
                            "current": self._mask_section(
                                section,
                                existing,
                            ),
                            "proposed": self._mask_section(
                                section,
                                proposed,
                            ),
                        }
                    )

        now = self._now()
        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            connection.execute("BEGIN")
            preview = repository.create_preview(
                resume_import_id=record["id"],
                preview_token=secrets.token_urlsafe(32),
                proposal=normalized,
                conflicts=conflicts,
                expires_at=self._iso(now + timedelta(hours=24)),
                now=self._iso(now),
            )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        return {
            **preview,
            "proposal": self._masked_proposal(normalized),
        }

    def confirm_resume_sections(
        self,
        preview_token: str,
        sections: list[str],
        *,
        corrections: dict | None = None,
    ) -> dict:
        if not sections or not set(sections) <= set(RESUME_SECTIONS):
            raise ResumeServiceError(
                "INVALID_INPUT",
                "sections must be a non-empty valid subset",
            )
        corrections = corrections or {}
        now = self._now()
        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            preview = repository.get_preview(preview_token)
            if preview is None:
                raise ResumeServiceError(
                    "PREVIEW_NOT_FOUND",
                    "resume preview not found",
                )
            if self._parse_iso(preview["expires_at"]) < now:
                raise ResumeServiceError(
                    "PREVIEW_EXPIRED",
                    "resume preview expired",
                )
            corrected = dict(preview["proposal"])
            corrected.update(corrections)
            try:
                validate_resume_proposal(corrected)
            except ResumeInputError as exc:
                raise ResumeServiceError(exc.code, str(exc)) from exc
            connection.execute("BEGIN")
            result = repository.confirm_sections(
                preview_token=preview_token,
                sections=sections,
                corrections=corrections,
                now=self._iso(now),
            )
            connection.commit()
            return result
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def get_candidate_profile(
        self,
        *,
        include_sensitive: bool = False,
    ) -> dict:
        profile = self._profile_or_none()
        if profile is None:
            raise ResumeServiceError(
                "PROFILE_NOT_FOUND",
                "candidate profile not found",
            )
        return profile if include_sensitive else mask_candidate_profile(profile)

    def get_profile_versions(self) -> list[dict]:
        connection = connect(self.config.database_path)
        try:
            return ResumeRepository(connection).list_profile_versions()
        finally:
            connection.close()

    def list_resume_documents(self) -> list[dict]:
        connection = connect(self.config.database_path)
        try:
            return ResumeRepository(connection).list_resume_documents()
        finally:
            connection.close()

    def get_resume_document(
        self,
        resume_id: str,
        *,
        include_sensitive: bool = False,
    ) -> dict:
        connection = connect(self.config.database_path)
        try:
            document = ResumeRepository(connection).get_resume_document(
                resume_id
            )
        finally:
            connection.close()
        if document is None:
            raise ResumeServiceError(
                "RESUME_NOT_FOUND",
                "confirmed resume not found",
            )
        return (
            document
            if include_sensitive
            else mask_candidate_profile(document)
        )

    def update_resume_document(
        self,
        resume_id: str,
        *,
        display_name: str,
        target_roles: list[str],
    ) -> dict:
        if not isinstance(display_name, str):
            raise ResumeServiceError(
                "INVALID_INPUT",
                "display_name must be a string",
            )
        normalized_name = display_name.strip()
        if not 1 <= len(normalized_name) <= 80:
            raise ResumeServiceError(
                "INVALID_INPUT",
                "display_name must contain 1 to 80 characters",
            )
        if not isinstance(target_roles, list) or not all(
            isinstance(role, str) for role in target_roles
        ):
            raise ResumeServiceError(
                "INVALID_INPUT",
                "target_roles must be a list of strings",
            )
        normalized_roles = list(
            dict.fromkeys(
                role.strip()
                for role in target_roles
                if role.strip()
            )
        )
        if len(normalized_roles) > 10:
            raise ResumeServiceError(
                "INVALID_INPUT",
                "target_roles must contain at most 10 labels",
            )
        if any(len(role) > 30 for role in normalized_roles):
            raise ResumeServiceError(
                "INVALID_INPUT",
                "target role labels must contain at most 30 characters",
            )

        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            connection.execute("BEGIN")
            document = repository.update_resume_document_metadata(
                resume_id,
                display_name=normalized_name,
                target_roles=normalized_roles,
                now=self._iso(self._now()),
            )
            if document is None:
                raise ResumeServiceError(
                    "RESUME_NOT_FOUND",
                    "confirmed resume not found",
                )
            connection.commit()
            return document
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def delete_resume_document(
        self,
        resume_id: str,
        *,
        confirm_name: str,
    ) -> dict:
        document = self.get_resume_document(
            resume_id,
            include_sensitive=True,
        )
        if confirm_name != document["display_name"]:
            raise ResumeServiceError(
                "CONFIRMATION_REQUIRED",
                "exact resume name confirmation is required",
            )

        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            connection.execute("BEGIN")
            deleted = repository.soft_delete_resume_document(
                resume_id,
                now=self._iso(self._now()),
            )
            if deleted is None:
                raise ResumeServiceError(
                    "RESUME_NOT_FOUND",
                    "confirmed resume not found",
                )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        original_removed = self._delete_private_path(
            deleted["private_path"]
        )
        return {
            "resume_id": resume_id,
            "original_removed": original_removed,
        }

    def get_resume_document_versions(self, resume_id: str) -> list[dict]:
        self.get_resume_document(resume_id)
        connection = connect(self.config.database_path)
        try:
            return ResumeRepository(
                connection
            ).list_resume_document_versions(resume_id)
        finally:
            connection.close()

    def get_resume_document_original(self, resume_id: str) -> dict:
        record, path = self._resume_original(resume_id)
        return {
            "filename": record["filename"],
            "source_kind": record["source_kind"],
            "size": path.stat().st_size,
            "preview_kind": (
                "html"
                if record["source_kind"] == "docx"
                else record["source_kind"]
            ),
        }

    def read_resume_document_original(self, resume_id: str) -> dict:
        record, path = self._resume_original(resume_id)
        return {
            "filename": record["filename"],
            "source_kind": record["source_kind"],
            "bytes": path.read_bytes(),
        }

    def render_resume_document_original(self, resume_id: str) -> dict:
        original = self.read_resume_document_original(resume_id)
        source_kind = original["source_kind"]
        if source_kind == "docx":
            return {
                **original,
                "bytes": self._docx_html(original["bytes"]),
                "media_type": "text/html; charset=utf-8",
            }
        return {
            **original,
            "media_type": (
                "application/pdf"
                if source_kind == "pdf"
                else "text/plain; charset=utf-8"
            ),
        }

    def get_profile_original(self) -> dict:
        record, path = self._profile_original()
        return {
            "filename": record["filename"],
            "source_kind": record["source_kind"],
            "size": path.stat().st_size,
            "preview_kind": (
                "html"
                if record["source_kind"] == "docx"
                else record["source_kind"]
            ),
        }

    def read_profile_original(self) -> dict:
        record, path = self._profile_original()
        return {
            "filename": record["filename"],
            "source_kind": record["source_kind"],
            "bytes": path.read_bytes(),
        }

    def render_profile_original(self) -> dict:
        original = self.read_profile_original()
        source_kind = original["source_kind"]
        if source_kind == "docx":
            return {
                **original,
                "bytes": self._docx_html(original["bytes"]),
                "media_type": "text/html; charset=utf-8",
            }
        return {
            **original,
            "media_type": (
                "application/pdf"
                if source_kind == "pdf"
                else "text/plain; charset=utf-8"
            ),
        }

    def discard_resume_import(self, preview_token: str) -> dict:
        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            connection.execute("BEGIN")
            try:
                result = repository.discard_preview_import(preview_token)
            except LookupError as exc:
                raise ResumeServiceError(
                    "PREVIEW_NOT_FOUND",
                    "resume preview not found",
                ) from exc
            except ValueError as exc:
                raise ResumeServiceError(
                    "CONFIRMATION_REQUIRED",
                    str(exc),
                ) from exc
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        self._delete_private_path(result.pop("private_path"))
        return result

    def delete_candidate_profile(self, confirm_text: str) -> dict:
        if confirm_text != "DELETE MY CAREER PROFILE":
            raise ResumeServiceError(
                "CONFIRMATION_REQUIRED",
                "exact deletion confirmation is required",
            )
        connection = connect(self.config.database_path)
        try:
            repository = ResumeRepository(connection)
            paths = repository.private_paths()
            connection.execute("BEGIN")
            counts = repository.delete_all()
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        for relative in paths:
            self._delete_private_path(relative)
        return counts

    def _valid_import(self, upload_token: str) -> dict:
        connection = connect(self.config.database_path)
        try:
            record = ResumeRepository(
                connection
            ).get_import_by_upload_token(upload_token)
        finally:
            connection.close()
        if record is None:
            raise ResumeServiceError(
                "UPLOAD_NOT_FOUND",
                "resume upload not found",
            )
        if self._parse_iso(record["expires_at"]) < self._now():
            raise ResumeServiceError(
                "UPLOAD_EXPIRED",
                "resume upload expired",
            )
        return record

    def _profile_or_none(self) -> dict | None:
        connection = connect(self.config.database_path)
        try:
            return ResumeRepository(connection).get_profile()
        finally:
            connection.close()

    def _profile_original(self) -> tuple[dict, Path]:
        connection = connect(self.config.database_path)
        try:
            record = ResumeRepository(connection).get_profile_original()
        finally:
            connection.close()
        if record is None:
            raise ResumeServiceError(
                "ORIGINAL_NOT_FOUND",
                "confirmed resume original not found",
            )
        return record, self._resolve_original_path(record)

    def _resume_original(self, resume_id: str) -> tuple[dict, Path]:
        connection = connect(self.config.database_path)
        try:
            record = ResumeRepository(connection).get_resume_original(
                resume_id
            )
        finally:
            connection.close()
        if record is None:
            raise ResumeServiceError(
                "RESUME_NOT_FOUND",
                "confirmed resume not found",
            )
        return record, self._resolve_original_path(record)

    @staticmethod
    def _default_ocr_service() -> ResumeOcrService:
        from ..ocr.rapidocr_engine import RapidOcrEngine

        return ResumeOcrService(
            engine_factory=RapidOcrEngine,
            renderer=PdfiumPageRenderer(),
        )

    def _publish_private_original(
        self,
        target: Path,
        raw_bytes: bytes,
    ) -> bool:
        self._assert_private_parent(target)
        target.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
        self._assert_private_parent(target)
        if target.is_symlink():
            raise OSError("private resume target cannot be a symlink")
        if target.is_file():
            return False
        if target.exists():
            raise OSError("private resume target is not a file")

        temporary = target.with_name(
            f".{target.name}.{secrets.token_hex(8)}.tmp"
        )
        descriptor: int | None = None
        published = False
        try:
            descriptor = os.open(
                temporary,
                os.O_WRONLY | os.O_CREAT | os.O_EXCL,
                0o600,
            )
            with os.fdopen(descriptor, "wb") as handle:
                descriptor = None
                handle.write(raw_bytes)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, target)
            published = True
            return True
        except Exception:
            if published:
                target.unlink(missing_ok=True)
            raise
        finally:
            if descriptor is not None:
                os.close(descriptor)
            temporary.unlink(missing_ok=True)

    def _assert_private_parent(self, target: Path) -> None:
        try:
            target.parent.resolve().relative_to(self.private_root)
        except ValueError as exc:
            raise OSError(
                "private resume target is outside private storage"
            ) from exc

    @staticmethod
    def _public_extraction(record: dict) -> dict[str, object]:
        metadata = record["extraction_metadata"]
        return {
            "extraction_method": str(record["extraction_method"]),
            "review_required": bool(metadata["review_required"]),
            "warnings": list(metadata["warning_codes"]),
            "pages": {
                "total": int(metadata["total_pages"]),
                "ocr": list(metadata["ocr_pages"]),
            },
        }

    def _resolve_original_path(self, record: dict) -> Path:
        path = (self.private_root / record["private_path"]).resolve()
        try:
            path.relative_to(self.private_root)
        except ValueError as exc:
            raise ResumeServiceError(
                "ORIGINAL_NOT_FOUND",
                "confirmed resume original is outside private storage",
            ) from exc
        if not path.is_file():
            raise ResumeServiceError(
                "ORIGINAL_NOT_FOUND",
                "confirmed resume original not found",
            )
        return path

    @staticmethod
    def _docx_html(raw_bytes: bytes) -> bytes:
        from docx import Document

        document = Document(io.BytesIO(raw_bytes))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = html.escape(paragraph.text.strip())
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = style.removeprefix("heading").strip() or "2"
                level = level if level in {"1", "2", "3", "4", "5", "6"} else "2"
                blocks.append(f"<h{level}>{text}</h{level}>")
            else:
                blocks.append(f"<p>{text}</p>")
        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = "".join(
                    f"<td>{html.escape(cell.text.strip())}</td>"
                    for cell in row.cells
                )
                rows.append(f"<tr>{cells}</tr>")
            blocks.append(f"<table>{''.join(rows)}</table>")
        body = "".join(blocks)
        return (
            "<article class=\"docx-preview\">"
            f"{body}"
            "</article>"
        ).encode("utf-8")

    @staticmethod
    def _masked_proposal(proposal: dict) -> dict:
        return mask_candidate_profile(
            {"sections": {key: proposal[key] for key in RESUME_SECTIONS}}
        )["sections"] | {
            "missing_fields": proposal["missing_fields"],
            "warnings": proposal["warnings"],
        }

    @staticmethod
    def _mask_section(section: str, value: object) -> object:
        if section != "basic":
            return value
        return mask_candidate_profile(
            {"sections": {"basic": value}}
        )["sections"]["basic"]

    def _now(self) -> datetime:
        value = self._clock()
        if value.tzinfo is None:
            raise ValueError("clock must be timezone-aware")
        return value.astimezone(timezone.utc)

    @staticmethod
    def _iso(value: datetime) -> str:
        return value.astimezone(timezone.utc).isoformat().replace(
            "+00:00",
            "Z",
        )

    @staticmethod
    def _parse_iso(value: str) -> datetime:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))

    def _delete_private_path(self, relative: str) -> bool:
        target = (self.private_root / relative).resolve()
        if self.private_root in target.parents:
            existed = target.is_file()
            target.unlink(missing_ok=True)
            return existed
        return False
