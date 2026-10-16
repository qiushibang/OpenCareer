import io
import sys
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

sys.path.insert(0, str(Path(__file__).parents[3]))

from extensions.career_mcp.domain.resume import (
    ResumeInputError,
    extract_resume_text,
    normalize_resume_text,
)
from extensions.career_mcp.services.resume_ocr import (
    MIN_NATIVE_PAGE_CHARS,
    NativePdfPage,
    ResumeOcrError,
    ResumeOcrResult,
    ResumeOcrService,
)


RESUME_LINES = [
    "Zhang San",
    "AI Engineer",
    "Python SQL local retrieval agent evaluation and deployment experience",
]
RESUME_TEXT = "\n".join(RESUME_LINES)


def _pdf_pages_bytes(page_lines: list[list[str]]) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for lines in page_lines:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        content = DecodedStreamObject()
        operations = ["BT", "/F1 12 Tf", "72 720 Td"]
        for index, line in enumerate(lines):
            if index:
                operations.append("0 -18 Td")
            operations.append(f"({line}) Tj")
        operations.append("ET")
        content.set_data("\n".join(operations).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
        page[NameObject("/MediaBox")] = ArrayObject(page["/MediaBox"])
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _pdf_bytes(lines: list[str], *, pages: int = 1) -> bytes:
    return _pdf_pages_bytes(
        [lines] + [["Continuation"] for _ in range(pages - 1)]
    )


def _encrypted_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes(lines: list[str]) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table Fact"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class FakeOcrService:
    def __init__(
        self,
        *,
        result: ResumeOcrResult,
        error: ResumeOcrError | None = None,
    ) -> None:
        self.result = result
        self.error = error
        self.calls: list[tuple[tuple[NativePdfPage, ...], bytes]] = []
        self.sparse_pages: list[int] = []

    def extract(
        self,
        pages: tuple[NativePdfPage, ...],
        *,
        raw_pdf: bytes,
    ) -> ResumeOcrResult:
        self.calls.append((pages, raw_pdf))
        self.sparse_pages = [
            page.page_number
            for page in pages
            if len(page.text) < MIN_NATIVE_PAGE_CHARS
        ]
        if self.error is not None:
            raise self.error
        return self.result


def _native_metadata(*, total_pages: int) -> dict[str, object]:
    return {
        "version": 1,
        "method": "native",
        "total_pages": total_pages,
        "native_pages": list(range(1, total_pages + 1)),
        "ocr_pages": [],
        "ocr_attempted_pages": [],
        "mean_confidence": None,
        "quality": "native",
        "review_required": False,
        "warning_codes": [],
        "engine": None,
    }


def test_text_pdf_and_docx_extract_equivalent_resume_lines():
    text = extract_resume_text(
        filename="resume.txt",
        mime="text/plain",
        text=RESUME_TEXT,
    )
    pdf = extract_resume_text(
        filename="resume.pdf",
        mime="application/pdf",
        raw_bytes=_pdf_bytes(RESUME_LINES),
    )
    docx = extract_resume_text(
        filename="resume.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        raw_bytes=_docx_bytes(RESUME_LINES),
    )

    assert text.text.splitlines() == RESUME_LINES
    assert pdf.text.splitlines() == RESUME_LINES
    assert docx.text.splitlines()[:3] == RESUME_LINES
    assert docx.text.splitlines()[-1] == "Table Fact"
    assert {text.source_kind, pdf.source_kind, docx.source_kind} == {
        "text",
        "pdf",
        "docx",
    }
    assert {
        text.extraction_method,
        pdf.extraction_method,
        docx.extraction_method,
    } == {"native"}
    assert text.extraction_metadata == _native_metadata(total_pages=1)
    assert pdf.extraction_metadata == _native_metadata(total_pages=1)
    assert docx.extraction_metadata == _native_metadata(total_pages=1)


def test_native_pdf_does_not_call_ocr_service() -> None:
    native_pdf = _pdf_bytes(RESUME_LINES)
    fake = FakeOcrService(
        result=ResumeOcrResult(
            text="not used",
            extraction_method="ocr",
            metadata={"engine": {"name": "not-used"}},
        )
    )

    extracted = extract_resume_text(
        filename="native.pdf",
        mime="application/pdf",
        raw_bytes=native_pdf,
        ocr_service=fake,
    )

    assert fake.calls == []
    assert extracted.extraction_method == "native"
    assert extracted.extraction_metadata["engine"] is None


def test_hybrid_pdf_passes_sparse_page_to_ocr_service() -> None:
    page_one = "Native resume content " + ("Python experience " * 6)
    raw_pdf = _pdf_pages_bytes([[page_one], []])
    metadata = {
        **_native_metadata(total_pages=2),
        "method": "hybrid",
        "native_pages": [1],
        "ocr_pages": [2],
        "ocr_attempted_pages": [2],
        "engine": {"name": "rapidocr-onnxruntime"},
    }
    fake = FakeOcrService(
        result=ResumeOcrResult(
            text=f"{page_one}\n\nScanned project experience",
            extraction_method="hybrid",
            metadata=metadata,
        )
    )

    extracted = extract_resume_text(
        filename="hybrid.pdf",
        mime="application/pdf",
        raw_bytes=raw_pdf,
        ocr_service=fake,
    )

    assert fake.sparse_pages == [2]
    assert len(fake.calls) == 1
    assert extracted.extraction_method == "hybrid"
    assert extracted.extraction_metadata["ocr_pages"] == [2]


def test_normalize_resume_text_normalizes_whitespace_without_merging_lines():
    assert normalize_resume_text("  Zhang\u00a0San \r\n\r\n AI\tEngineer  ") == (
        "Zhang San\n\nAI Engineer"
    )


@pytest.mark.parametrize(
    ("kwargs", "code"),
    [
        (
            {
                "filename": "resume.txt",
                "mime": "text/plain",
                "text": "too short",
            },
            "INVALID_INPUT",
        ),
        (
            {
                "filename": "resume.pdf",
                "mime": "application/pdf",
                "raw_bytes": _pdf_bytes([], pages=1),
            },
            "NO_EXTRACTABLE_TEXT",
        ),
        (
            {
                "filename": "resume.docx",
                "mime": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                "raw_bytes": b"not a docx",
            },
            "INVALID_INPUT",
        ),
        (
            {
                "filename": "resume.rtf",
                "mime": "application/rtf",
                "raw_bytes": b"resume content",
            },
            "UNSUPPORTED_FORMAT",
        ),
    ],
)
def test_resume_extraction_rejects_invalid_inputs(kwargs, code):
    with pytest.raises(ResumeInputError) as error:
        extract_resume_text(**kwargs)

    assert error.value.code == code


def test_resume_extraction_rejects_pdf_over_page_limit():
    with pytest.raises(ResumeInputError) as error:
        extract_resume_text(
            filename="resume.pdf",
            mime="application/pdf",
            raw_bytes=_pdf_bytes(["Resume"], pages=21),
        )

    assert error.value.code == "PDF_PAGE_LIMIT"


@pytest.mark.parametrize(
    ("raw_bytes", "code"),
    [
        (_encrypted_pdf_bytes(), "INVALID_INPUT"),
        (b"not a PDF", "INVALID_INPUT"),
        (b"x" * (10 * 1024 * 1024 + 1), "FILE_TOO_LARGE"),
    ],
)
def test_pdf_safety_errors_use_stable_codes(
    raw_bytes: bytes,
    code: str,
) -> None:
    with pytest.raises(ResumeInputError) as error:
        extract_resume_text(
            filename="resume.pdf",
            mime="application/pdf",
            raw_bytes=raw_bytes,
        )

    assert error.value.code == code


def test_pdf_requiring_more_than_ten_ocr_pages_is_rejected() -> None:
    engine_factory_called = False

    def engine_factory() -> Any:
        nonlocal engine_factory_called
        engine_factory_called = True
        raise AssertionError("engine must not initialize")

    service = ResumeOcrService(
        engine_factory=engine_factory,
        renderer=object(),  # type: ignore[arg-type]
    )

    with pytest.raises(ResumeInputError) as error:
        extract_resume_text(
            filename="resume.pdf",
            mime="application/pdf",
            raw_bytes=_pdf_pages_bytes([[] for _ in range(11)]),
            ocr_service=service,
        )

    assert error.value.code == "OCR_PAGE_LIMIT"
    assert engine_factory_called is False


def test_pdf_final_short_text_does_not_leak_extracted_content() -> None:
    private_text = "private short"
    fake = FakeOcrService(
        result=ResumeOcrResult(
            text=private_text,
            extraction_method="ocr",
            metadata={
                **_native_metadata(total_pages=1),
                "method": "ocr",
                "native_pages": [],
                "ocr_pages": [1],
            },
        )
    )

    with pytest.raises(ResumeInputError) as error:
        extract_resume_text(
            filename="resume.pdf",
            mime="application/pdf",
            raw_bytes=_pdf_pages_bytes([[]]),
            ocr_service=fake,
        )

    assert error.value.code == "NO_EXTRACTABLE_TEXT"
    assert private_text not in str(error.value)


def test_resume_extraction_sanitizes_original_filename():
    extracted = extract_resume_text(
        filename="../../private/resume.txt",
        mime="text/plain",
        text=RESUME_TEXT,
    )

    assert extracted.filename == "resume.txt"
