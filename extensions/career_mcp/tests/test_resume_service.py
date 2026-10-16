import sys
import sqlite3
import stat
from io import BytesIO
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

sys.path.insert(0, str(Path(__file__).parents[3]))

from extensions.career_mcp.config import CareerConfig
from extensions.career_mcp.database.migrations import connect, migrate
from extensions.career_mcp.services.resume_ocr import (
    NativePdfPage,
    ResumeOcrError,
    ResumeOcrResult,
)
from extensions.career_mcp.services import resumes as resume_service_module
from extensions.career_mcp.services.resumes import (
    ResumeService,
    ResumeServiceError,
)


NOW = datetime(2026, 8, 6, 8, 0, tzinfo=timezone.utc)
TEXT = """Zhang San
Phone: 13800001234
Email: zhang@example.com
AI Engineer with Python experience
"""
PROPOSAL = {
    "basic": {
        "name": "Zhang San",
        "phone": "13800001234",
        "email": "zhang@example.com",
        "location": "Beijing",
        "summary": "AI Engineer",
    },
    "education": [],
    "experience": [],
    "project": [],
    "skill": [
        {"name": "Python", "level": None, "evidence": None}
    ],
    "missing_fields": [],
    "warnings": [],
}


def _service(tmp_path, **kwargs: Any):
    config = CareerConfig(data_dir=tmp_path / "career")
    return ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: NOW,
        **kwargs,
    )


def _pdf_bytes(page_texts: list[str | None]) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        if text is None:
            continue
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(
            f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        )
        page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _blank_pdf_bytes(*, pages: int = 1) -> bytes:
    return _pdf_bytes([None] * pages)


class FixedOcrService:
    def __init__(
        self,
        *,
        result: ResumeOcrResult | None = None,
        error: ResumeOcrError | None = None,
    ) -> None:
        self.result = result
        self.error = error
        self.calls: list[tuple[tuple[NativePdfPage, ...], bytes]] = []

    def extract(
        self,
        pages: tuple[NativePdfPage, ...],
        *,
        raw_pdf: bytes,
    ) -> ResumeOcrResult:
        self.calls.append((pages, raw_pdf))
        if self.error is not None:
            raise self.error
        assert self.result is not None
        return self.result


def _extraction_metadata(
    *,
    method: str,
    total_pages: int,
    native_pages: list[int],
    ocr_pages: list[int],
    warnings: list[str] | None = None,
) -> dict[str, object]:
    warning_codes = warnings or []
    return {
        "version": 1,
        "method": method,
        "total_pages": total_pages,
        "native_pages": native_pages,
        "ocr_pages": ocr_pages,
        "ocr_attempted_pages": ocr_pages,
        "mean_confidence": 0.91 if ocr_pages else None,
        "quality": "review" if warning_codes else (
            "good" if ocr_pages else "native"
        ),
        "review_required": bool(warning_codes),
        "warning_codes": warning_codes,
        "engine": (
            {
                "name": "rapidocr-onnxruntime",
                "version": "3.9.2",
                "model_ids": ["det", "cls", "rec"],
                "model_sha256": ["a" * 64, "b" * 64, "c" * 64],
            }
            if ocr_pages
            else None
        ),
    }


def _resume_import_count(config: CareerConfig) -> int:
    connection = connect(config.database_path)
    try:
        migrate(connection)
        return int(
            connection.execute(
                "SELECT COUNT(*) FROM resume_imports"
            ).fetchone()[0]
        )
    finally:
        connection.close()


def test_stage_resume_is_private_idempotent_and_omits_text(tmp_path):
    service = _service(tmp_path)

    first = service.stage_resume(
        filename="../../resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    second = service.stage_resume(
        filename="copy.txt",
        mime="text/plain",
        text=TEXT,
    )

    assert first["upload_token"] == second["upload_token"]
    assert first["filename"] == "resume.txt"
    assert first["source_kind"] == "text"
    assert first["characters"] == len(first["sha256"]) * 0 + len(TEXT.strip())
    assert first["extraction_method"] == "native"
    assert first["review_required"] is False
    assert first["warnings"] == []
    assert first["pages"] == {"total": 1, "ocr": []}
    assert "text" not in first
    assert "private_path" not in first
    originals = list((tmp_path / "private" / "resumes").iterdir())
    assert len(originals) == 1


def test_resume_preview_does_not_write_profile_and_masks_sensitive_fields(
    tmp_path,
):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )

    imported = service.get_resume_import_text(staged["upload_token"])
    preview = service.preview_resume_import(
        staged["upload_token"],
        PROPOSAL,
    )

    assert imported["text"] == TEXT.strip()
    assert preview["proposal"]["basic"]["phone"] == "*******1234"
    with pytest.raises(ResumeServiceError) as error:
        service.get_candidate_profile()
    assert error.value.code == "PROFILE_NOT_FOUND"


def test_confirm_sections_masks_default_profile_and_versions_changes(
    tmp_path,
):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(
        staged["upload_token"],
        PROPOSAL,
    )

    first = service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    second = service.confirm_resume_sections(
        preview["preview_token"],
        ["skill"],
        corrections={
            "skill": [
                {
                    "name": "Python",
                    "level": "advanced",
                    "evidence": None,
                }
            ]
        },
    )
    masked = service.get_candidate_profile()
    clear = service.get_candidate_profile(include_sensitive=True)

    assert first["version"] == 1
    assert second["version"] == 2
    assert masked["sections"]["basic"]["name"] == "Z*** S**"
    assert masked["sections"]["basic"]["phone"] == "*******1234"
    assert masked["sections"]["basic"]["email"] == "z***@example.com"
    assert clear["sections"]["basic"]["phone"] == "13800001234"
    assert clear["sections"]["skill"][0]["level"] == "advanced"


def test_new_preview_reports_section_conflict_without_overwriting(tmp_path):
    service = _service(tmp_path)
    first = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(first["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    changed = {
        **PROPOSAL,
        "basic": {**PROPOSAL["basic"], "location": "Shanghai"},
    }

    second_preview = service.preview_resume_import(
        first["upload_token"],
        changed,
    )
    profile = service.get_candidate_profile(include_sensitive=True)

    assert second_preview["conflicts"][0]["section"] == "basic"
    assert profile["sections"]["basic"]["location"] == "Beijing"


def test_delete_requires_phrase_and_removes_private_originals(tmp_path):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )

    with pytest.raises(ResumeServiceError) as error:
        service.delete_candidate_profile("delete")
    assert error.value.code == "CONFIRMATION_REQUIRED"

    deleted = service.delete_candidate_profile(
        "DELETE MY CAREER PROFILE"
    )

    assert deleted["profiles"] == 1
    assert not list((tmp_path / "private" / "resumes").glob("*"))


def test_delete_profile_does_not_require_retired_generation_tables(tmp_path):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    connection = connect(service.config.database_path)
    connection.execute("PRAGMA foreign_keys = OFF")
    connection.execute("DROP TABLE tailored_resume_exports")
    connection.execute("DROP TABLE tailored_resume_drafts")
    connection.execute("DROP TABLE job_resume_matches")
    connection.commit()
    connection.close()

    deleted = service.delete_candidate_profile(
        "DELETE MY CAREER PROFILE"
    )

    assert deleted["profiles"] == 1
    assert "matches" not in deleted


def test_invalid_or_expired_tokens_do_not_expose_text(tmp_path):
    service = _service(tmp_path)

    with pytest.raises(ResumeServiceError) as error:
        service.get_resume_import_text("missing")

    assert error.value.code == "UPLOAD_NOT_FOUND"
    assert TEXT not in str(error.value)


def test_reupload_after_expiry_refreshes_token_without_duplicate_original(
    tmp_path,
):
    current = [NOW]
    config = CareerConfig(data_dir=tmp_path / "career")
    service = ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: current[0],
    )
    first = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    current[0] = NOW + timedelta(hours=25)

    second = service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )

    assert second["upload_token"] != first["upload_token"]
    assert service.get_resume_import_text(second["upload_token"])["text"] == (
        TEXT.strip()
    )
    assert len(list((tmp_path / "private" / "resumes").iterdir())) == 1


def test_profile_page_reads_versions_and_original_without_private_path(tmp_path):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="candidate-resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )

    versions = service.get_profile_versions()
    original = service.get_profile_original()
    content = service.read_profile_original()

    assert versions[0]["version"] == 1
    assert versions[0]["confirmed_sections"] == ["basic"]
    assert original == {
        "filename": "candidate-resume.txt",
        "source_kind": "text",
        "size": len(TEXT.strip().encode()),
        "preview_kind": "text",
    }
    assert content["filename"] == "candidate-resume.txt"
    assert content["source_kind"] == "text"
    assert content["bytes"] == TEXT.strip().encode()
    assert "private_path" not in original


def test_docx_profile_original_renders_escaped_read_only_html(tmp_path):
    stream = BytesIO()
    document = Document()
    document.add_heading("Zhang <script>alert(1)</script>", level=1)
    document.add_paragraph(
        "AI Engineer with enough resume content for safe local preview."
    )
    document.save(stream)
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="candidate.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        raw_bytes=stream.getvalue(),
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )

    rendered = service.render_profile_original()

    assert rendered["media_type"] == "text/html; charset=utf-8"
    assert b"&lt;script&gt;alert(1)&lt;/script&gt;" in rendered["bytes"]
    assert b"<script>" not in rendered["bytes"]


def test_resume_library_reads_each_confirmed_import_independently(tmp_path):
    service = _service(tmp_path)
    for index, filename in enumerate(
        ("通用简历.txt", "产品经理简历.txt"),
        start=1,
    ):
        staged = service.stage_resume(
            filename=filename,
            mime="text/plain",
            text=f"{TEXT}\nTailored resume {index}",
        )
        proposal = {
            **PROPOSAL,
            "basic": {
                **PROPOSAL["basic"],
                "summary": f"Tailored resume {index}",
            },
        }
        preview = service.preview_resume_import(
            staged["upload_token"],
            proposal,
        )
        service.confirm_resume_sections(
            preview["preview_token"],
            ["basic"],
            corrections={},
        )

    documents = service.list_resume_documents()
    selected_id = next(
        item["id"]
        for item in documents
        if item["filename"] == "产品经理简历.txt"
    )
    masked = service.get_resume_document(selected_id)
    clear = service.get_resume_document(
        selected_id,
        include_sensitive=True,
    )
    versions = service.get_resume_document_versions(selected_id)
    original = service.get_resume_document_original(selected_id)

    assert {item["filename"] for item in documents} == {
        "产品经理简历.txt",
        "通用简历.txt",
    }
    assert masked["sections"]["basic"]["name"] == "Z*** S**"
    assert clear["sections"]["basic"]["summary"] == "Tailored resume 2"
    assert versions
    assert original["filename"] == "产品经理简历.txt"
    assert "private_path" not in original

    with pytest.raises(ResumeServiceError) as error:
        service.get_resume_document("missing")
    assert error.value.code == "RESUME_NOT_FOUND"


def test_resume_document_metadata_is_normalized_and_validated(tmp_path):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="通用简历.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    resume_id = service.list_resume_documents()[0]["id"]

    updated = service.update_resume_document(
        resume_id,
        display_name=" AI 产品经理简历 ",
        target_roles=[
            "AI 产品经理",
            " AI 产品经理 ",
            "Agent 产品经理",
            "",
        ],
    )

    assert updated["display_name"] == "AI 产品经理简历"
    assert updated["target_roles"] == [
        "AI 产品经理",
        "Agent 产品经理",
    ]
    assert service.get_resume_document(
        resume_id,
        include_sensitive=True,
    )["display_name"] == "AI 产品经理简历"

    invalid_values = [
        {"display_name": "", "target_roles": []},
        {"display_name": "x" * 81, "target_roles": []},
        {
            "display_name": "简历",
            "target_roles": [str(index) for index in range(11)],
        },
        {"display_name": "简历", "target_roles": ["x" * 31]},
    ]
    for payload in invalid_values:
        with pytest.raises(ResumeServiceError) as error:
            service.update_resume_document(resume_id, **payload)
        assert error.value.code == "INVALID_INPUT"


def test_single_resume_delete_requires_name_and_preserves_shared_profile(
    tmp_path,
):
    service = _service(tmp_path)
    resume_ids = []
    for index, (filename, sections) in enumerate(
        (
            ("通用简历.txt", ["basic"]),
            ("产品经理简历.txt", ["skill"]),
        ),
        start=1,
    ):
        staged = service.stage_resume(
            filename=filename,
            mime="text/plain",
            text=f"{TEXT}\nTailored resume {index}",
        )
        preview = service.preview_resume_import(
            staged["upload_token"],
            PROPOSAL,
        )
        service.confirm_resume_sections(
            preview["preview_token"],
            sections,
            corrections={},
        )
        resume_ids.append(
            next(
                item["id"]
                for item in service.list_resume_documents()
                if item["filename"] == filename
            )
        )

    with pytest.raises(ResumeServiceError) as error:
        service.delete_resume_document(
            resume_ids[0],
            confirm_name="错误名称",
        )
    assert error.value.code == "CONFIRMATION_REQUIRED"

    deleted = service.delete_resume_document(
        resume_ids[0],
        confirm_name="通用简历.txt",
    )

    assert deleted == {
        "resume_id": resume_ids[0],
        "original_removed": True,
    }
    assert [item["id"] for item in service.list_resume_documents()] == [
        resume_ids[1]
    ]
    assert service.get_candidate_profile(
        include_sensitive=True,
    )["sections"]["basic"]["name"] == "Zhang San"
    assert len(service.get_profile_versions()) == 2
    assert len(list((tmp_path / "private" / "resumes").glob("*"))) == 1
    with pytest.raises(ResumeServiceError) as error:
        service.get_resume_document(resume_ids[0])
    assert error.value.code == "RESUME_NOT_FOUND"


def test_deleted_resume_can_be_reimported_from_the_same_file(tmp_path):
    service = _service(tmp_path)
    staged = service.stage_resume(
        filename="通用简历.txt",
        mime="text/plain",
        text=TEXT,
    )
    preview = service.preview_resume_import(staged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    resume = service.list_resume_documents()[0]
    service.delete_resume_document(
        resume["id"],
        confirm_name=resume["display_name"],
    )

    restaged = service.stage_resume(
        filename="重新上传简历.txt",
        mime="text/plain",
        text=TEXT,
    )

    assert restaged["upload_token"] != staged["upload_token"]
    assert restaged["filename"] == "重新上传简历.txt"
    assert restaged["characters"] == len(TEXT.strip())
    preview = service.preview_resume_import(restaged["upload_token"], PROPOSAL)
    service.confirm_resume_sections(
        preview["preview_token"],
        ["basic"],
        corrections={},
    )
    assert service.list_resume_documents()[0]["display_name"] == (
        "重新上传简历.txt"
    )


@pytest.mark.parametrize(
    (
        "filename",
        "pages",
        "result",
        "expected_method",
        "expected_review",
        "expected_warnings",
        "expected_ocr_pages",
    ),
    [
        (
            "scanned.pdf",
            1,
            ResumeOcrResult(
                text="OCR recovered candidate resume content",
                extraction_method="ocr",
                metadata=_extraction_metadata(
                    method="ocr",
                    total_pages=1,
                    native_pages=[],
                    ocr_pages=[1],
                ),
            ),
            "ocr",
            False,
            [],
            [1],
        ),
        (
            "hybrid.pdf",
            2,
            ResumeOcrResult(
                text=(
                    "Native candidate profile with sufficient content "
                    "for the first page and OCR project experience"
                ),
                extraction_method="hybrid",
                metadata=_extraction_metadata(
                    method="hybrid",
                    total_pages=2,
                    native_pages=[1],
                    ocr_pages=[2],
                ),
            ),
            "hybrid",
            False,
            [],
            [2],
        ),
        (
            "review.pdf",
            1,
            ResumeOcrResult(
                text="OCR recovered low confidence resume content",
                extraction_method="ocr",
                metadata=_extraction_metadata(
                    method="ocr",
                    total_pages=1,
                    native_pages=[],
                    ocr_pages=[1],
                    warnings=["OCR_LOW_CONFIDENCE"],
                ),
            ),
            "ocr",
            True,
            ["OCR_LOW_CONFIDENCE"],
            [1],
        ),
    ],
)
def test_stage_pdf_returns_safe_extraction_provenance(
    filename: str,
    pages: int,
    result: ResumeOcrResult,
    expected_method: str,
    expected_review: bool,
    expected_warnings: list[str],
    expected_ocr_pages: list[int],
    tmp_path: Path,
) -> None:
    ocr = FixedOcrService(result=result)
    service = _service(tmp_path, ocr_service=ocr)
    raw_pdf = (
        _pdf_bytes(
            [
                (
                    "Native candidate profile with Python SQL agent "
                    "evaluation deployment and product experience"
                ),
                None,
            ]
        )
        if expected_method == "hybrid"
        else _blank_pdf_bytes(pages=pages)
    )

    staged = service.stage_resume(
        filename=filename,
        mime="application/pdf",
        raw_bytes=raw_pdf,
    )
    imported = service.get_resume_import_text(staged["upload_token"])

    assert staged["extraction_method"] == expected_method
    assert staged["review_required"] is expected_review
    assert staged["warnings"] == expected_warnings
    assert staged["pages"] == {
        "total": pages,
        "ocr": expected_ocr_pages,
    }
    assert imported["extraction_method"] == expected_method
    assert imported["extraction_metadata"] == result.metadata
    assert imported["review_required"] is expected_review
    assert imported["warnings"] == expected_warnings
    assert "private_path" not in imported
    assert len(ocr.calls) == 1


def test_native_pdf_does_not_construct_ocr_service(tmp_path: Path) -> None:
    factory_called = False

    def factory() -> FixedOcrService:
        nonlocal factory_called
        factory_called = True
        raise AssertionError("native PDF must not construct OCR")

    service = _service(tmp_path, ocr_service_factory=factory)

    staged = service.stage_resume(
        filename="native.pdf",
        mime="application/pdf",
        raw_bytes=_pdf_bytes(
            [
                (
                    "Native resume content with Python SQL agent evaluation "
                    "deployment and product experience exceeding eighty chars"
                )
            ]
        ),
    )

    assert staged["extraction_method"] == "native"
    assert factory_called is False


def test_runtime_unavailable_with_native_text_returns_review_warning(
    tmp_path: Path,
) -> None:
    metadata = _extraction_metadata(
        method="native",
        total_pages=1,
        native_pages=[1],
        ocr_pages=[],
        warnings=["OCR_RUNTIME_UNAVAILABLE"],
    )
    metadata["ocr_attempted_pages"] = [1]
    ocr = FixedOcrService(
        result=ResumeOcrResult(
            text="Native text remains sufficient for safe user review",
            extraction_method="native",
            metadata=metadata,
        )
    )
    service = _service(tmp_path, ocr_service=ocr)

    staged = service.stage_resume(
        filename="sparse.pdf",
        mime="application/pdf",
        raw_bytes=_pdf_bytes(
            ["Native text remains sufficient for safe user review"]
        ),
    )

    assert staged["extraction_method"] == "native"
    assert staged["review_required"] is True
    assert staged["warnings"] == ["OCR_RUNTIME_UNAVAILABLE"]


@pytest.mark.parametrize(
    ("failure", "expected_code"),
    [
        (
            ResumeOcrError(
                "OCR_FAILED",
                "private recognized text must not escape",
            ),
            "OCR_FAILED",
        ),
        (
            ResumeOcrError(
                "OCR_UNAVAILABLE",
                "/private/model/path must not escape",
            ),
            "OCR_UNAVAILABLE",
        ),
    ],
)
def test_ocr_failure_leaves_no_private_file_or_database_row(
    failure: ResumeOcrError,
    expected_code: str,
    tmp_path: Path,
) -> None:
    config = CareerConfig(data_dir=tmp_path / "career")
    factory_calls = 0

    def factory() -> FixedOcrService:
        nonlocal factory_calls
        factory_calls += 1
        if expected_code == "OCR_UNAVAILABLE":
            raise failure
        return FixedOcrService(error=failure)

    service = ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: NOW,
        ocr_service_factory=factory,
    )

    with pytest.raises(ResumeServiceError) as error:
        service.stage_resume(
            filename="scanned.pdf",
            mime="application/pdf",
            raw_bytes=_blank_pdf_bytes(),
        )

    assert error.value.code == expected_code
    assert "private" not in str(error.value)
    assert not list((tmp_path / "private" / "resumes").glob("*"))
    assert _resume_import_count(config) == 0
    assert factory_calls == 1


def test_repository_failure_removes_newly_published_original(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = CareerConfig(data_dir=tmp_path / "career")

    class FailingRepository:
        def __init__(self, connection: sqlite3.Connection) -> None:
            migrate(connection)

        def create_or_reuse_import(self, **kwargs: Any) -> dict:
            raise sqlite3.OperationalError("private SQL failure")

    monkeypatch.setattr(
        resume_service_module,
        "ResumeRepository",
        FailingRepository,
    )
    service = ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: NOW,
    )

    with pytest.raises(ResumeServiceError) as error:
        service.stage_resume(
            filename="resume.txt",
            mime="text/plain",
            text=TEXT,
        )

    assert error.value.code == "PERSISTENCE_ERROR"
    assert str(tmp_path) not in str(error.value)
    assert "SQL failure" not in str(error.value)
    assert not list((tmp_path / "private" / "resumes").glob("*"))
    assert _resume_import_count(config) == 0


def test_private_rename_failure_leaves_no_file_or_database_row(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = CareerConfig(data_dir=tmp_path / "career")
    service = ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: NOW,
    )

    def fail_replace(source: object, target: object) -> None:
        raise OSError("private rename failure")

    monkeypatch.setattr(resume_service_module.os, "replace", fail_replace)

    with pytest.raises(ResumeServiceError) as error:
        service.stage_resume(
            filename="resume.txt",
            mime="text/plain",
            text=TEXT,
        )

    assert error.value.code == "PERSISTENCE_ERROR"
    assert str(tmp_path) not in str(error.value)
    assert "rename failure" not in str(error.value)
    assert not list((tmp_path / "private" / "resumes").glob("*"))
    assert _resume_import_count(config) == 0


def test_database_failure_does_not_delete_preexisting_original(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = CareerConfig(data_dir=tmp_path / "career")
    service = ResumeService(
        config,
        private_root=tmp_path / "private",
        clock=lambda: NOW,
    )
    service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )
    original = next((tmp_path / "private" / "resumes").iterdir())

    class FailingRepository:
        def __init__(self, connection: sqlite3.Connection) -> None:
            migrate(connection)

        def create_or_reuse_import(self, **kwargs: Any) -> dict:
            raise sqlite3.OperationalError("injected failure")

    monkeypatch.setattr(
        resume_service_module,
        "ResumeRepository",
        FailingRepository,
    )

    with pytest.raises(ResumeServiceError) as error:
        service.stage_resume(
            filename="copy.txt",
            mime="text/plain",
            text=TEXT,
        )

    assert error.value.code == "PERSISTENCE_ERROR"
    assert original.is_file()
    assert _resume_import_count(config) == 1


def test_private_resume_directory_symlink_cannot_escape_storage_root(
    tmp_path: Path,
) -> None:
    config = CareerConfig(data_dir=tmp_path / "career")
    private_root = tmp_path / "private"
    outside = tmp_path / "outside"
    private_root.mkdir()
    outside.mkdir()
    (private_root / "resumes").symlink_to(outside, target_is_directory=True)
    service = ResumeService(
        config,
        private_root=private_root,
        clock=lambda: NOW,
    )

    with pytest.raises(ResumeServiceError) as error:
        service.stage_resume(
            filename="resume.txt",
            mime="text/plain",
            text=TEXT,
        )

    assert error.value.code == "PERSISTENCE_ERROR"
    assert list(outside.iterdir()) == []
    assert _resume_import_count(config) == 0


def test_new_private_resume_original_is_owner_only(tmp_path: Path) -> None:
    service = _service(tmp_path)

    service.stage_resume(
        filename="resume.txt",
        mime="text/plain",
        text=TEXT,
    )

    original = next((tmp_path / "private" / "resumes").iterdir())
    assert stat.S_IMODE(original.stat().st_mode) == 0o600
